"""Generate the Dino Coach user manual as Word (DOCX) and PDF from one source.

Requires python-docx, Pillow and a Chromium binary (Playwright's bundled Chromium works).
Run from the repository root:
    python3 scripts/generate-dino-manual.py
Set CHROME to override the Chromium path. Outputs go to public/documents/.
Colours are the official club palette from tailwind.config.ts.
"""
from pathlib import Path
import base64
import html
import io
import os
import subprocess
import tempfile

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from PIL import Image

STEM = '20260928-Dino-Coach-User-Manual-Rev00'
DATE = '28 September 2026'
REVISION = 'Rev00'
RULES_VERSION = '2026-27-rev07'
OUT = Path('public/documents')
LOGO = Path('public/images/logo.jpg')
LOGO_ALT = 'Newcomb and District Cricket Club logo'
CHROME = os.environ.get('CHROME', '/opt/pw-browsers/chromium-1194/chrome-linux/chrome')

MAROON, GOLD, SKY, NAVY, CREAM = '880000', 'EDC266', '8CC6D1', '162845', 'FBF7F0'
GOLD_TINT, SKY_TINT, MAROON_TINT = 'FBF0D9', 'E8F4F6', 'FCE4E4'
# One accent per chapter keeps the document lively without inventing colours.
ACCENTS = [MAROON, NAVY, MAROON, NAVY, MAROON, NAVY, MAROON, NAVY, MAROON]

INTRO = ('Build a squad of NDCC cricketers, choose their fantasy roles and earn points from published match '
         'performances. This guide explains registration, squad selection, your live wallet, buying from and '
         'selling back to the player pool, scoring and competition rules.')

# Block types: h1, h2, p, steps, table, tip, note, links.
CONTENT = [
    ('h1', 'At a glance'),
    ('table', ['Topic', 'Current rule'], [
        ['Entry', 'AUD 25.00; managers must be 18 or older'],
        ['Starting budget', '15,000,000 virtual Dino Dollars'],
        ['Squad', '15 different players: 11 starters and 4 bench players'],
        ['Leadership', 'One captain and one vice-captain; both score double'],
        ['Sales', 'Original purchase cost refunded'],
        ['Player pool', 'Buy from and sell back to the pool; no inter-team trades'],
        ['Weekly window', 'Tuesday 00:00 to before Saturday 11:00, Melbourne time'],
    ]),
    ('h2', 'Start here'),
    ('p', 'Register and confirm your email, complete your manager profile and pay your entry fee. Fill all 15 slots, '
          'choose your captain and vice-captain, then select Submit squad before the round locks.'),
    ('p', 'Existing teams retain their recorded purchase costs. Read the current rules and accept them directly in the '
          'squad builder, keeping your selected players in place. You can also accept them in My account.'),
    ('tip', 'Open Dino Coach', 'www.ndcc.com.au/fantasy'),

    ('h1', 'Register and get ready to play'),
    ('h2', 'Create your account'),
    ('steps', [
        'Open Dino Coach and choose Register / Login. If you already have an NDCC account, sign in with it.',
        'Enter your display name, team name, email address, password and date of birth. You must be at least 18. '
        'Read and accept the current rules.',
        'Select Register and follow the confirmation link sent to your email. Check spam or junk if it does not '
        'arrive. The registration screen has a Resend confirmation email option.',
        'Return to My account. Complete your manager profile if prompted and select Save profile. Use your '
        'registration email.',
    ]),
    ('h2', 'Team names and payment'),
    ('p', 'Choose a name suitable for a community cricket club. Flagged names need committee approval before payment; '
          'the league manager may replace and lock an unsuitable name.'),
    ('p', 'When eligible, select Pay AUD 25.00 entry and complete checkout. Return to My account and check for Entry '
          'paid - you can pick your team. A pending payment is still awaiting confirmation. Keep your payment '
          'reference and report a suspected payment issue before paying again.'),
    ('p', 'Club-approved complimentary entries do not require payment. Demo access also bypasses payment, but demo '
          'teams are not eligible for prizes.'),
    ('h2', 'Finish your first squad on time'),
    ('p', 'Your registration does not expire if your first squad is incomplete. Save a draft as you build your team, '
          'then fill all 15 slots and submit before the round locks. Normal competition and transfer rules still apply.'),
    ('note', 'Something not saving?', 'If a technical issue prevents saving, use the Dino Coach feedback form. Include '
             'your team name, registration email and the error shown beside the save buttons.'),
    ('h2', 'When the rules change'),
    ('p', 'Read the current rules using the link in the squad builder. Tick the acceptance box and select Accept rules '
          'and keep my selections. Your selected players stay in place. Then save or submit your squad.'),

    ('h1', 'Build your squad and compare players'),
    ('p', 'Choose the correct season and open My squad. Select 15 different eligible players within 15,000,000 Dino '
          'Dollars. A player can appear only once in your team; other managers can select the same cricketer.'),
    ('table', ['Fantasy role', 'Playing XI', 'Bench', 'Total'], [
        ['BAT - batter', '4', '1', '5'],
        ['AR - all-rounder', '2', '1', '3'],
        ['WK - wicket keeper', '1', '1', '2'],
        ['BOWL - bowler', '4', '1', '5'],
        ['All roles', '11', '4', '15'],
    ]),
    ('p', 'Any eligible player can fill any fantasy slot. Their listed cricket role is a guide; the assigned slot '
          'determines scoring. Junior-only players are excluded. Juniors who also play senior cricket need '
          'league-manager eligibility confirmation.'),
    ('h2', 'Pick and save'),
    ('steps', [
        'Search the catalogue or sort by name or price. Select a player in the keyboard/touch control, then choose '
        'Assign selected player on a slot. On a computer, you can also drag a player into a slot.',
        'Fill all playing and bench slots. Sell / remove clears a selection in your unsaved draft. The wallet preview '
        'changes immediately.',
        'Choose Captain for one starter and Vice for a different starter. Both score double and neither can be on '
        'the bench.',
        'Check 15/15 slots, an affordable balance and both leadership choices. Select Submit squad and wait for '
        'confirmation. Reopen My squad to check the saved team.',
    ]),
    ('tip', 'Drafts', 'Save draft keeps incomplete work. After your first full squad, edits must be saved within the '
            'transfer window and before the round deadline. A changed squad in another session requires a reload '
            'before saving.'),
    ('h2', 'Read the player cards'),
    ('p', 'Cards show batting runs, bowling wickets and fielding catches, with match records, stumpings, run-outs and '
          'maidens where available. The source period appears on each card. Historical totals are used until '
          'published records exist for the selected season. Compare like periods; Not recorded means unknown, not '
          'zero. Cards contain recorded totals rather than invented ratings.'),

    ('h1', 'How points are awarded'),
    ('p', 'Your assigned fantasy role determines points for runs, wickets and catches. The table applies to starters. '
          'Bench players earn zero.'),
    ('table', ['Event', 'BAT', 'AR', 'WK', 'BOWL'], [
        ['Each run', '1.75', '1.5', '1.5', '1'],
        ['Each wicket', '10', '15', '10', '20'],
        ['Each catch', '10', '10', '15', '10'],
    ]),
    ('table', ['Event or milestone', 'Points for any playing role'], [
        ['Each run-out', '10'],
        ['Each stumping', '10'],
        ['Each maiden over', '5'],
        ['Not-out innings', '10'],
        ['50-99 runs in an innings', '20 bonus'],
        ['100 or more runs in an innings', '50 bonus'],
        ['5-6 wickets in an innings', '25 bonus'],
        ['7 or more wickets in an innings', '50 bonus'],
    ]),
    ('p', 'Only the highest batting milestone and highest bowling milestone apply in an innings. A century earns the '
          '50-point bonus without another 20 for fifty. Seven wickets earn the 50-point bonus without another 25 for five.'),
    ('p', 'Captain and vice-captain each receive double their full assigned-role points, including bonuses. There is '
          'no duck penalty and no player-of-the-match bonus.'),
    ('h2', 'A scoring example'),
    ('tip', 'Worked example', 'A BAT starter makes 60 runs, takes one catch and finishes not out. They earn 105 for '
            'runs, 10 for the catch, 10 for not out and 20 for reaching fifty: 145 points. As captain or '
            'vice-captain, they earn 290 points.'),

    ('h1', 'Your wallet and player sales'),
    ('p', 'Every team starts with 15,000,000 virtual Dino Dollars. The Team wallet shows the starting budget, saved '
          'spending and saved money available. My squad also shows the balance after your unsaved selections.'),
    ('p', 'Selecting or removing players changes the preview immediately. Save draft or Submit squad confirms those '
          'edits. Market transactions update the saved wallet after they succeed. The live connection receives '
          'changes from another session, with automatic refresh during reconnection. If an error says the balance '
          'may be out of date, reconnect and check it before acting.'),
    ('table', ['Action', 'Money available'], [
        ['Buy a player', 'Decreases by the current published purchase price'],
        ['Sell a player', "Increases by that player's original purchase cost"],
        ['Price review', "Unchanged; the player's market value may change"],
        ['Unsaved squad edit', 'Preview only until the save succeeds'],
    ]),
    ('h2', 'Example in Dino Dollars'),
    ('table', ['Step', 'Available money'], [
        ['Starting balance', '15,000,000'],
        ['Buy a player for 800,000', '14,200,000'],
        ['Their market price rises to 1,200,000', '14,200,000'],
        ['Sell that player for the original 800,000 cost', '15,000,000'],
    ]),
    ('h2', 'Sell or buy separately'),
    ('p', 'Open Transfers. Choose Player to sell and check the displayed refund, then select Sell back to pool. To fill '
          'an empty slot, choose Player to buy, select the empty slot and choose Buy into empty slot. A sale or '
          'separate purchase leaves your squad as a draft. Return to My squad, fill all slots, check leadership and '
          'submit before the deadline.'),
    ('h2', 'Replace a player in one action'),
    ('p', 'Choose both Player to sell and Player to buy. Check the after-replacement balance, then select Sell and buy '
          'replacement. The new player inherits the outgoing slot and leadership selection. Check those choices in My '
          'squad afterwards. A failed transaction does not charge your wallet.'),

    ('h1', 'Using the shared player pool'),
    ('p', 'All player purchases and sales take place through the shared player pool. Transfers between teams are '
          'unavailable. You do not need another manager to buy your player before you can sell them.'),
    ('h2', 'Sell back to the pool'),
    ('steps', [
        'Open Transfers during the open window and choose Player to sell.',
        "Check the refund shown. Selling returns the original purchase cost, even if the player's current published "
        'value has changed.',
        'Select Sell back to pool and wait for confirmation. Your saved money available increases after the sale succeeds.',
        'Fill the empty slot from the pool and submit your complete squad before the deadline.',
    ]),
    ('h2', 'Buying and replacing'),
    ('p', 'Choose an eligible player you do not already own. A purchase charges the current published price. To '
          'replace a player in one step, choose both players and select Sell and buy replacement.'),
    ('p', "The same cricketer can appear in several managers' squads. Another team selecting a player does not remove "
          'that player from the pool or prevent your team from selecting them.'),
    ('table', ['Action', 'Effect'], [
        ['Sell back to pool', 'Refunds original purchase cost; leaves an empty slot'],
        ['Buy from pool', 'Charges current price; fills an empty slot'],
        ['Sell and buy replacement', 'Refund and purchase complete together'],
        ['Another team selects a player', 'Player remains available in the shared pool'],
    ]),
    ('p', 'Earlier pending inter-team offers are cancelled. They do not move players or money. Existing squads and '
          'recorded purchase costs are preserved.'),
    ('note', 'Keep enough money', 'Keep enough money for all 15 slots. Check the saved wallet after each completed '
             'transaction and check your captain and vice-captain before submitting.'),

    ('h1', 'Transfer windows and player prices'),
    ('table', ['Melbourne time', 'Availability'], [
        ['Tuesday 00:00', 'Window opens'],
        ['Tuesday to Friday', 'Open, subject to competition controls'],
        ['Saturday before 11:00', 'Open until the closing time'],
        ['Saturday 11:00 to Tuesday before 00:00', 'Closed'],
    ]),
    ('p', 'Transfers, purchases and sales are unlimited and free, with no points penalty. All times use '
          'Australia/Melbourne, including daylight saving. The website checks the window when you save; beginning '
          'before the deadline does not reserve a transaction. Round locks and season controls may close changes earlier.'),
    ('h2', 'Opening prices'),
    ('p', 'Opening prices use the supplied season statistics. The highest-ranked player starts at 2,000,000 Dino '
          'Dollars and players without historical statistics start at 100,000. Opening prices round to the nearest '
          '1,000 Dino Dollars. Later price reviews and manual changes round upwards to whole 1,000. External player '
          'cards identify their source period and missing figures.'),
    ('h2', 'Automatic reviews'),
    ('p', 'Reviews occur after every two regular rounds: rounds 2, 4, 6 and so on. A review becomes eligible on the '
          'following Monday at 09:00 Melbourne time and runs at the next daily pricing check once published results '
          'are available.'),
    ('p', 'Strong performances can raise prices and weaker performances can lower them. Prices stay between 100,000 '
          'and 2,000,000 Dino Dollars. Before a player or grade starts, unplayed weeks do not count as zero '
          'appearances or lower the performance average or price. Finals do not change prices. Late results enter a '
          'later review. Completed reviews are not applied twice. The league manager may make recorded price or '
          'eligibility corrections; automatic changes resume at the next review.'),
    ('h2', 'Purchase cost and market value'),
    ('p', "Purchase cost is what your wallet paid for a player and what a sale refunds. Market value is the player's "
          'current published value, used for squad-value comparisons and standings ties. A price rise does not add '
          'cash to the wallet. Retaining a player while saving your squad preserves their purchase cost.'),

    ('h1', 'Standings and competition rules'),
    ('table', ['Standings', 'What they show'], [
        ['Player Standings', 'Real NDCC cricketers ranked by published base performance points, without '
                             'assigned-slot or captain bonuses.'],
        ['Manager Standings', 'Fantasy teams ranked by their playing XI points, including assigned roles and both '
                              'leadership bonuses. Bench players contribute zero.'],
    ]),
    ('p', 'The same cricketer can show a different score in Player Standings and your fantasy team. Check the '
          'selected season and published round before comparing totals.'),
    ('h2', 'Ties and prizes'),
    ('p', 'Manager ties are resolved by total points, then current squad market value, then team name alphabetically. '
          'The round-robin leader prize is 300 Dino Dollars. The committee publishes the label and description for '
          'the highest total squad-value prize. Dino Dollars are virtual; no real-money value is implied. Demo teams '
          'are not eligible for prizes.'),
    ('h2', 'Published results and corrections'),
    ('p', 'Points depend on published match results and may not appear immediately after a match. Report the player, '
          'round or match, the score shown and what needs checking. Include a source or match reference where possible.'),
    ('p', 'The 2026/2027 competition is a pilot. The league manager may make a reasonable adjustment to protect '
          'fairness or operation when a scoring defect, data issue, technical fault or unintended rules outcome is '
          'found. Material changes must be communicated and recorded. Check the Rules page and club notices.'),
    ('h2', 'Fair play and account care'),
    ('p', 'Use a suitable team name, give accurate registration details and keep your password private. Treat '
          'players, managers and volunteers respectfully. Report suspected abuse or unexpected behaviour through the '
          'feedback form.'),
    ('p', 'For withdrawal, entry payment help or restoration of a removed team, contact the club through the feedback '
          'form. Do not create another paid entry to resolve an existing account issue.'),
    ('h2', 'A note from the club'),
    ('tip', 'Thank you for playing', 'This is our first time running Dino Coach, so please be patient while we work '
            'through any issues. If something is not working, or you have a suggestion, use the feedback form. Saj '
            'and Rick will review it and look into what can be improved.'),

    ('h1', 'Help and useful links'),
    ('table', ['Problem', 'What to do'], [
        ['No confirmation email', 'Check spam or junk, then use Resend confirmation email. Verified NDCC accounts '
                                  'should sign in.'],
        ['Forgotten password', 'Use Forgot password on the sign-in page and follow the reset email.'],
        ['Payment pending or team locked', 'Check My account and your payment reference. Report the issue before '
                                           'paying again.'],
        ['Squad will not submit', 'Check 15 unique players, all slots, budget, leadership, eligibility, accepted '
                                  'rules and open selection.'],
        ['Pool transaction rejected', 'Check the window, squad, published price and available money. Reload if '
                                      'another session changed your squad.'],
        ['Balance seems stale', 'Check the wallet connection message, reconnect and reload. Unsaved edits are only '
                                'a preview.'],
        ['Account disabled or team removed', 'Request reactivation or restoration through the feedback form.'],
        ['Scores or statistics look wrong', 'Check the source period and published round, then provide player and '
                                            'match details.'],
    ]),
    ('h2', 'Send an issue or suggestion'),
    ('p', 'Use the feedback form with your name, reply email and a description. Include the page, team name, '
          'approximate time and error message. Do not include passwords or card details. You do not need to sign in. '
          'Select Send feedback and keep the displayed reference. Feedback is private and reviewed by Saj and Rick.'),
    ('h2', 'Keep Dino Coach handy'),
    ('p', "On iPhone or iPad, use Safari: Share, then Add to Home Screen. On Android, use Chrome's menu and choose "
          'Install app or Add to Home screen if offered. On a computer, use the browser install option or bookmark '
          'the page. Internet access is required.'),
    ('h2', 'Useful links'),
    ('links', [
        ('Home', 'www.ndcc.com.au/fantasy'),
        ('Account', 'www.ndcc.com.au/fantasy/account'),
        ('Rules', 'www.ndcc.com.au/fantasy/rules'),
        ('Feedback', 'www.ndcc.com.au/fantasy/feedback'),
    ]),
]

def logo_bytes():
    # The site logo is 1184 px wide; 480 px is ample for 3.4 cm and keeps both files small.
    image = Image.open(LOGO).convert('RGB')
    image.thumbnail((480, 480))
    buffer = io.BytesIO()
    image.save(buffer, 'JPEG', quality=85, optimize=True)
    return buffer.getvalue()


HEADER_TEXT = f'NDCC / DINO COACH 2026/2027    {DATE} | {REVISION} | Rules {RULES_VERSION}'


# ---------------------------------------------------------------- Word (DOCX)
def shade(cell, fill):
    props = cell._tc.get_or_add_tcPr()
    element = OxmlElement('w:shd')
    element.set(qn('w:val'), 'clear')
    element.set(qn('w:color'), 'auto')
    element.set(qn('w:fill'), fill)
    props.append(element)


def cell_borders(cell, colour, size=8, sides=('top', 'left', 'bottom', 'right')):
    props = cell._tc.get_or_add_tcPr()
    borders = OxmlElement('w:tcBorders')
    for side in ('top', 'left', 'bottom', 'right'):
        element = OxmlElement(f'w:{side}')
        if side in sides:
            element.set(qn('w:val'), 'single')
            element.set(qn('w:sz'), str(size))
            element.set(qn('w:color'), colour)
        else:
            element.set(qn('w:val'), 'nil')
        borders.append(element)
    props.append(borders)


def cell_margins(cell, top=90, bottom=90, left=140, right=140):
    props = cell._tc.get_or_add_tcPr()
    margins = OxmlElement('w:tcMar')
    for side, value in (('top', top), ('bottom', bottom), ('left', left), ('right', right)):
        element = OxmlElement(f'w:{side}')
        element.set(qn('w:w'), str(value))
        element.set(qn('w:type'), 'dxa')
        margins.append(element)
    props.append(margins)


def run(paragraph, text, bold=False, colour=NAVY, size=10.5):
    item = paragraph.add_run(text)
    item.bold = bold
    item.font.size = Pt(size)
    item.font.color.rgb = RGBColor.from_string(colour)
    return item


def spacing(paragraph, before=0, after=6):
    paragraph.paragraph_format.space_before = Pt(before)
    paragraph.paragraph_format.space_after = Pt(after)


def box(document, title, text, fill, edge, label):
    table = document.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.rows[0].cells[0]
    shade(cell, fill)
    cell_borders(cell, edge, size=24, sides=('left',))
    cell_margins(cell, 120, 120, 220, 160)
    first = cell.paragraphs[0]
    spacing(first, 0, 2)
    run(first, f'{label}  ', bold=True, colour=MAROON if label == 'IMPORTANT' else NAVY, size=8.5)
    run(first, title, bold=True, size=11)
    body = cell.add_paragraph()
    spacing(body, 0, 0)
    run(body, text, size=10.5)
    spacing(document.add_paragraph(), 0, 2)


def data_table(document, headers, rows, accent):
    table = document.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for index, header in enumerate(headers):
        cell = table.rows[0].cells[index]
        shade(cell, accent)
        cell_borders(cell, accent, sides=('bottom',))
        cell_margins(cell)
        paragraph = cell.paragraphs[0]
        spacing(paragraph, 0, 0)
        run(paragraph, header.upper(), bold=True, colour='FFFFFF', size=9)
    for number, row in enumerate(rows):
        cells = table.add_row().cells
        total = row[0] == 'All roles'
        for index, value in enumerate(row):
            fill = GOLD_TINT if total else (CREAM if number % 2 == 0 else 'FFFFFF')
            shade(cells[index], fill)
            cell_borders(cells[index], 'E6DCCB', size=4, sides=('bottom',))
            cell_margins(cells[index])
            paragraph = cells[index].paragraphs[0]
            spacing(paragraph, 0, 0)
            run(paragraph, value, bold=index == 0 or total, size=10)
    spacing(document.add_paragraph(), 0, 4)


def page_field(paragraph):
    for kind, text in (('begin', None), (None, 'PAGE'), ('end', None)):
        item = paragraph.add_run()
        item.font.size = Pt(8.5)
        item.font.color.rgb = RGBColor.from_string(NAVY)
        if kind:
            element = OxmlElement('w:fldChar')
            element.set(qn('w:fldCharType'), kind)
        else:
            element = OxmlElement('w:instrText')
            element.set(qn('xml:space'), 'preserve')
            element.text = text
        item._r.append(element)


TC_ORDER = ['cnfStyle', 'tcW', 'gridSpan', 'hMerge', 'vMerge', 'tcBorders', 'shd', 'noWrap', 'tcMar',
            'textDirection', 'tcFitText', 'vAlign', 'hideMark']
R_ORDER = ['rStyle', 'rFonts', 'b', 'bCs', 'i', 'iCs', 'caps', 'smallCaps', 'strike', 'dstrike', 'outline', 'shadow',
           'emboss', 'imprint', 'noProof', 'snapToGrid', 'vanish', 'webHidden', 'color', 'spacing', 'w', 'kern',
           'position', 'sz', 'szCs', 'highlight', 'u', 'effect', 'bdr', 'shd', 'fitText', 'vertAlign', 'rtl', 'cs',
           'em', 'lang', 'eastAsianLayout', 'specVanish', 'oMath']


def schema_order(document):
    # Word rejects out-of-order properties, so sort what the helpers appended.
    for tag, order in (('w:tcPr', TC_ORDER), ('w:rPr', R_ORDER)):
        for props in document.element.body.iter(qn(tag)):
            children = list(props)
            rank = lambda child: order.index(child.tag.split('}')[1]) if child.tag.split('}')[1] in order else len(order)
            for child in children:
                props.remove(child)
            for child in sorted(children, key=rank):
                props.append(child)


def build_docx(path):
    document = Document()
    section = document.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.left_margin = section.right_margin = Cm(2)
    section.top_margin, section.bottom_margin = Cm(2), Cm(1.8)
    normal = document.styles['Normal']
    normal.font.name = 'Arial'
    normal.element.rPr.rFonts.set(qn('w:eastAsia'), 'Arial')
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor.from_string(NAVY)
    for name, size, colour in (('Title', 30, MAROON), ('Heading 1', 18, MAROON), ('Heading 2', 13, NAVY)):
        style = document.styles[name]
        style.font.name = 'Arial'
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(colour)

    header = section.header.paragraphs[0]
    run(header, HEADER_TEXT, bold=True, colour=MAROON, size=8)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run(footer, 'Newcomb and District Cricket Club  |  Page ', colour=NAVY, size=8.5)
    page_field(footer)

    # Cover band: maroon block with gold rule, then the logo and introduction.
    band = document.add_table(rows=1, cols=2)
    band.alignment = WD_TABLE_ALIGNMENT.CENTER
    logo_cell, title_cell = band.rows[0].cells
    logo_cell.width, title_cell.width = Cm(4.2), Cm(12.8)
    for cell in (logo_cell, title_cell):
        shade(cell, MAROON)
        cell_borders(cell, GOLD, size=36, sides=('bottom',))
        cell_margins(cell, 260, 260, 260, 200)
    logo = logo_cell.paragraphs[0].add_run().add_picture(io.BytesIO(logo_bytes()), width=Cm(3.4))
    logo._inline.docPr.set('descr', LOGO_ALT)
    logo._inline.docPr.set('title', LOGO_ALT)
    title = title_cell.paragraphs[0]
    spacing(title, 0, 2)
    run(title, 'Dino Coach user manual', bold=True, colour='FFFFFF', size=26)
    subtitle = title_cell.add_paragraph()
    spacing(subtitle, 0, 2)
    run(subtitle, 'Rules and manager guide | 2026/2027', bold=True, colour=GOLD, size=13)
    club = title_cell.add_paragraph()
    spacing(club, 0, 0)
    run(club, f'Newcomb and District Cricket Club  |  {DATE}', colour='FFFFFF', size=10)
    spacing(document.add_paragraph(), 0, 6)
    intro = document.add_paragraph()
    spacing(intro, 0, 10)
    run(intro, INTRO, size=11.5)

    chapter = 0
    for block in CONTENT:
        kind = block[0]
        if kind == 'h1':
            chapter += 1
            if chapter > 1:
                document.add_page_break()
            accent = ACCENTS[chapter - 1]
            heading = document.add_paragraph(style='Heading 1')
            spacing(heading, 4, 8)
            badge = heading.add_run(f' {chapter} ')
            badge.font.color.rgb = RGBColor.from_string(GOLD if accent == NAVY else 'FFFFFF')
            badge.font.highlight_color = None
            properties = badge._r.get_or_add_rPr()
            background = OxmlElement('w:shd')
            background.set(qn('w:val'), 'clear')
            background.set(qn('w:fill'), accent)
            properties.append(background)
            text = heading.add_run(f'  {block[1]}')
            text.font.color.rgb = RGBColor.from_string(accent)
        elif kind == 'h2':
            heading = document.add_paragraph(style='Heading 2')
            spacing(heading, 10, 4)
            marker = heading.add_run('■ ')
            marker.font.color.rgb = RGBColor.from_string(GOLD)
            heading.add_run(block[1])
        elif kind == 'p':
            paragraph = document.add_paragraph()
            spacing(paragraph, 0, 6)
            run(paragraph, block[1])
        elif kind == 'steps':
            for number, text in enumerate(block[1], 1):
                paragraph = document.add_paragraph()
                paragraph.paragraph_format.left_indent = Cm(0.9)
                paragraph.paragraph_format.first_line_indent = Cm(-0.9)
                spacing(paragraph, 0, 5)
                run(paragraph, f'{number}   ', bold=True, colour=MAROON, size=11)
                run(paragraph, text)
        elif kind == 'table':
            data_table(document, block[1], block[2], ACCENTS[chapter - 1] if chapter else MAROON)
        elif kind == 'tip':
            box(document, block[1], block[2], SKY_TINT, SKY, 'TIP')
        elif kind == 'note':
            box(document, block[1], block[2], GOLD_TINT, GOLD, 'IMPORTANT')
        elif kind == 'links':
            data_table(document, ['Page', 'Address'], [list(item) for item in block[1]], NAVY)
    schema_order(document)
    document.core_properties.title = 'Dino Coach user manual'
    document.core_properties.subject = f'Rules and manager guide 2026/2027, {REVISION}'
    document.core_properties.author = 'Newcomb and District Cricket Club'
    document.save(path)


# ------------------------------------------------------------------ PDF (HTML)
def build_html():
    logo = base64.b64encode(logo_bytes()).decode()
    e = html.escape
    parts = []
    chapter = 0
    for block in CONTENT:
        kind = block[0]
        if kind == 'h1':
            chapter += 1
            accent = ACCENTS[chapter - 1]
            if chapter > 1:
                parts.append('</section>')
            parts.append(f'<section class="chapter" style="--accent:#{accent}">'
                         f'<h1><span class="badge">{chapter}</span>{e(block[1])}</h1>')
        elif kind == 'h2':
            parts.append(f'<h2>{e(block[1])}</h2>')
        elif kind == 'p':
            parts.append(f'<p>{e(block[1])}</p>')
        elif kind == 'steps':
            parts.append('<ol class="steps">' + ''.join(f'<li>{e(item)}</li>' for item in block[1]) + '</ol>')
        elif kind in ('table', 'links'):
            headers, rows = (block[1], block[2]) if kind == 'table' else (['Page', 'Address'], block[1])
            head = ''.join(f'<th>{e(item)}</th>' for item in headers)
            body = ''.join(('<tr class="total">' if row[0] == 'All roles' else '<tr>')
                           + ''.join(f'<td>{e(value)}</td>' for value in row) + '</tr>' for row in rows)
            parts.append(f'<table><thead><tr>{head}</tr></thead><tbody>{body}</tbody></table>')
        elif kind in ('tip', 'note'):
            label = 'Tip' if kind == 'tip' else 'Important'
            parts.append(f'<aside class="{kind}"><p class="label">{label}</p><p class="title">{e(block[1])}</p>'
                         f'<p>{e(block[2])}</p></aside>')
    parts.append('</section>')
    return f'''<!doctype html><html lang="en-AU"><head><meta charset="utf-8"><title>Dino Coach user manual</title>
<style>
@page {{ size: A4; margin: 20mm 18mm 18mm;
  @top-left {{ content: "{e(HEADER_TEXT)}"; font: 700 7.5pt Arial, sans-serif; color: #{MAROON}; }}
  @bottom-right {{ content: "Newcomb and District Cricket Club  |  Page " counter(page); font: 8pt Arial, sans-serif; color: #{NAVY}; }} }}
* {{ box-sizing: border-box; }}
body {{ margin: 0; font: 10.2pt/1.5 Arial, "Liberation Sans", sans-serif; color: #{NAVY}; }}
.cover {{ display: flex; gap: 18px; align-items: center; background: #{MAROON}; color: #fff; padding: 22px 24px;
  border-radius: 14px; border-bottom: 6px solid #{GOLD}; }}
.cover img {{ width: 112px; border-radius: 8px; background: #fff; }}
.cover h1 {{ margin: 0 0 4px; font-size: 28pt; line-height: 1.1; }}
.cover .sub {{ margin: 0 0 4px; color: #{GOLD}; font-weight: 700; font-size: 13pt; }}
.cover .meta {{ margin: 0; font-size: 10pt; }}
.stripe {{ height: 6px; margin: 10px 0 14px; border-radius: 3px;
  background: linear-gradient(90deg, #{MAROON} 0 34%, #{GOLD} 34% 67%, #{SKY} 67% 100%); }}
.intro {{ font-size: 11.5pt; }}
.chapter {{ break-before: page; }}
.chapter:first-of-type {{ break-before: auto; }}
.chapter h1 {{ display: flex; align-items: center; gap: 10px; margin: 6px 0 10px; font-size: 18pt; color: var(--accent);
  border-bottom: 3px solid #{GOLD}; padding-bottom: 6px; }}
.badge {{ display: inline-flex; width: 30px; height: 30px; align-items: center; justify-content: center; border-radius: 50%;
  background: var(--accent); color: #fff; font-size: 13pt; }}
h2 {{ margin: 16px 0 6px; font-size: 12.5pt; color: #{NAVY}; break-after: avoid; }}
h2::before {{ content: ""; display: inline-block; width: 9px; height: 9px; margin-right: 8px; background: #{GOLD}; border-radius: 2px; }}
p {{ margin: 0 0 7px; }}
table {{ width: 100%; border-collapse: separate; border-spacing: 0; margin: 6px 0 12px; border-radius: 10px; overflow: hidden;
  break-inside: avoid; border: 1px solid #E6DCCB; }}
th {{ background: var(--accent, #{MAROON}); color: #fff; text-align: left; font-size: 8.5pt; letter-spacing: .04em;
  text-transform: uppercase; padding: 7px 10px; }}
td {{ padding: 6px 10px; border-top: 1px solid #EFE6D6; vertical-align: top; }}
td:first-child {{ font-weight: 700; }}
tbody tr:nth-child(odd) td {{ background: #{CREAM}; }}
tr.total td {{ background: #{GOLD_TINT} !important; font-weight: 700; }}
.steps {{ list-style: none; counter-reset: step; padding: 0; margin: 4px 0 10px; }}
.steps li {{ counter-increment: step; position: relative; padding: 2px 0 6px 38px; break-inside: avoid; }}
.steps li::before {{ content: counter(step); position: absolute; left: 0; top: 0; width: 26px; height: 26px; border-radius: 50%;
  background: #{MAROON}; color: #fff; font-weight: 700; display: flex; align-items: center; justify-content: center; }}
aside {{ margin: 8px 0 12px; padding: 10px 14px; border-radius: 10px; break-inside: avoid; }}
aside.tip {{ background: #{SKY_TINT}; border-left: 6px solid #{SKY}; }}
aside.note {{ background: #{GOLD_TINT}; border-left: 6px solid #{GOLD}; }}
aside p {{ margin: 0; }}
aside .label {{ font-size: 7.5pt; font-weight: 700; letter-spacing: .08em; text-transform: uppercase; color: #{MAROON}; }}
aside .title {{ font-weight: 700; font-size: 11pt; margin-bottom: 2px; }}
</style></head><body>
<header class="cover"><img src="data:image/jpeg;base64,{logo}" alt="{e(LOGO_ALT)}">
<div><h1>Dino Coach user manual</h1><p class="sub">Rules and manager guide | 2026/2027</p>
<p class="meta">Newcomb and District Cricket Club | {DATE}</p></div></header>
<div class="stripe"></div>
<p class="intro">{e(INTRO)}</p>
{''.join(parts)}
</body></html>'''


def main():
    OUT.mkdir(parents=True, exist_ok=True)
    build_docx(OUT / f'{STEM}.docx')
    with tempfile.TemporaryDirectory() as folder:
        page = Path(folder) / 'manual.html'
        page.write_text(build_html(), encoding='utf-8')
        subprocess.run([CHROME, '--headless', '--no-sandbox', '--disable-gpu', '--no-pdf-header-footer',
                        f'--print-to-pdf={(OUT / f"{STEM}.pdf").resolve()}', page.as_uri()],
                       check=True, capture_output=True)
    print(f'Wrote {OUT / STEM}.docx and .pdf')


if __name__ == '__main__':
    main()
